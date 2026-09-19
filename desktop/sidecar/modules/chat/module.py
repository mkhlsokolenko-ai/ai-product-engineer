"""Модуль «Чат» — тред №1 оболочки.

Возможности MVP: треды/темы + история (локальный SQLite), выбор профиля (code/ask/standard),
скиллы из UI (подмешиваются в system), вложения → RAG (rag_index/rag_search через шлюз),
память треда (последние реплики в контексте), мультиагенты (передача задачи по ролям).

Всё в модели — только через gateway (единый JWT/квота). Добавление фич = правка ЭТОГО модуля,
не ядра. Позже методики скиллов можно грузить из skills/*/SKILL.md (сейчас — краткие подсказки).
"""
from __future__ import annotations

import json
import re
import urllib.error
import urllib.request
from pathlib import Path

from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from ... import auth, config, db, gateway

MANIFEST = {"id": "chat", "title": "Чат", "icon": "chat", "ui": "chat", "order": 10}

router = APIRouter()

# Краткие подсказки скиллов (v1). Позже — загрузка полной методики из skills/<id>/SKILL.md.
SKILLS = {
    "email-draft": "Пиши деловые письма по структуре: цель, контекст, просьба, дедлайн.",
    "icp-interviewer": "Помогай раскрыть портрет клиента: триггер, костыль, успех, готовность платить.",
    "devils-advocate": "Жёстко критикуй идею: без похвал, каждое возражение с аргументом.",
    "idea-scorer": "Оценивай идею по рубрике: боль, данные, выполнимость, экономика.",
    "finance-report": "Собирай финсводку: только числа из источника, без выдумок.",
}
BASE_SYSTEM = "Ты — рабочий ИИ-ассистент. Отвечай по делу, без воды. Если данных нет — скажи прямо, не выдумывай."
# Пресеты ролей для мультиагентов — пользователь выбирает, кого подключить под задачу.
ROLE_PRESETS = {
    "researcher": ("Ресёрчер", "Собери факты и контекст по задаче. Только проверяемое, помечай неуверенность."),
    "analyst": ("Аналитик", "На основе фактов найди закономерности и выводы. Структурируй."),
    "critic": ("Критик", "Проверь выводы на прочность: где натяжки, чего не хватает, что перепроверить."),
    "writer": ("Редактор", "Собери итог в чистый структурированный ответ для делового читателя."),
    "planner": ("Планировщик", "Разложи задачу на шаги с ответственными и сроками."),
    "finance": ("Финансист", "Посчитай экономику: затраты, эффект, риски. Числа только из данных."),
}
DEFAULT_ROLES = ["researcher", "analyst", "critic"]


class ThreadIn(BaseModel):
    title: str = "Новый чат"
    profile: str = "standard"
    skills: list[str] = []
    favorite: int = 0


class SendIn(BaseModel):
    prompt: str


class AttachIn(BaseModel):
    name: str = "документ"
    documents: list[str]


class AgentsIn(BaseModel):
    task: str
    roles: list[str] = []       # id из ROLE_PRESETS (быстрые пресеты)
    agent_ids: list[int] = []   # id из каталога (модуль agents) — приоритетнее roles


class ExportIn(BaseModel):
    format: str = "md"      # md | docx | xlsx (pdf делает Electron через printToPDF)


def _downloads() -> Path:
    d = Path.home() / "Downloads"
    return d if d.exists() else Path.home()


def _safe(name: str) -> str:
    return re.sub(r"[^\w\-. ]", "_", name or "").strip()[:60] or "chat"


def _sid(thread_id: int) -> str:
    return f"desktop-thread-{thread_id}"


def _system_for(skills: list[str]) -> str:
    hints = [f"[{s}] {SKILLS[s]}" for s in skills if s in SKILLS]
    return BASE_SYSTEM + ("\n\nАктивные методики:\n" + "\n".join(hints) if hints else "")


def _history(thread_id: int, limit: int = 6) -> str:
    rows = db.q("SELECT role,content FROM messages WHERE thread_id=? ORDER BY id DESC LIMIT ?",
                (thread_id, limit))
    rows = list(reversed(rows))
    return "\n".join(f"{'Ты' if r['role']=='user' else 'Ассистент'}: {r['content']}" for r in rows)


# ── скиллы для UI ──
@router.get("/skills")
def skills() -> list[dict]:
    return [{"id": k, "hint": v} for k, v in SKILLS.items()]


# ── треды ──
@router.get("/threads")
def list_threads() -> list[dict]:
    rows = db.q("SELECT id,title,profile,skills,favorite,updated_at FROM threads "
                "ORDER BY favorite DESC, updated_at DESC")
    for r in rows:
        r["skills"] = [s for s in (r["skills"] or "").split(",") if s]
    return rows


@router.post("/threads")
def create_thread(body: ThreadIn) -> dict:
    ts = db.now()
    tid = db.run("INSERT INTO threads(title,profile,skills,created_at,updated_at) VALUES(?,?,?,?,?)",
                 (body.title, body.profile, ",".join(body.skills), ts, ts))
    return {"id": tid, "title": body.title, "profile": body.profile, "skills": body.skills}


@router.delete("/threads/{thread_id}")
def delete_thread(thread_id: int) -> dict:
    db.run("DELETE FROM messages WHERE thread_id=?", (thread_id,))
    db.run("DELETE FROM threads WHERE id=?", (thread_id,))
    return {"ok": True}


@router.get("/threads/{thread_id}/messages")
def messages(thread_id: int) -> list[dict]:
    rows = db.q("SELECT id,role,content,meta,created_at FROM messages WHERE thread_id=? ORDER BY id",
                (thread_id,))
    for r in rows:
        r["meta"] = json.loads(r["meta"] or "{}")
    return rows


@router.patch("/threads/{thread_id}")
def update_thread(thread_id: int, body: ThreadIn) -> dict:
    db.run("UPDATE threads SET title=?,profile=?,skills=?,favorite=?,updated_at=? WHERE id=?",
           (body.title, body.profile, ",".join(body.skills), int(body.favorite), db.now(), thread_id))
    return {"ok": True}


@router.post("/threads/{thread_id}/autotitle")
def autotitle(thread_id: int) -> dict:
    """Авто-название темы по первым репликам (короткий вызов модели)."""
    msgs = db.q("SELECT content FROM messages WHERE thread_id=? AND role='user' ORDER BY id LIMIT 2",
                (thread_id,))
    if not msgs:
        return {"ok": False, "error": "empty"}
    seed = "\n".join(m["content"] for m in msgs)[:800]
    try:
        r = gateway.call("chat", {
            "prompt": f"Придумай короткое название темы чата (3-5 слов, без кавычек и точки) по началу диалога:\n{seed}",
            "session_id": _sid(thread_id) + "-title", "profile": "standard",
            "system": "Верни ТОЛЬКО название, без пояснений.", "max_tokens": 30})
    except (gateway.AuthRequired, gateway.GatewayError):
        return {"ok": False, "error": "gateway"}
    title = (r.get("text", "") or "").strip().strip('"').splitlines()[0][:60] or "Новый чат"
    db.run("UPDATE threads SET title=?,updated_at=? WHERE id=?", (title, db.now(), thread_id))
    return {"ok": True, "title": title}


# ── вложения → RAG (S3/Qdrant через шлюз) ──
@router.post("/threads/{thread_id}/attach")
def attach(thread_id: int, body: AttachIn) -> dict:
    try:
        res = gateway.call("rag_index", {"documents": body.documents, "session_id": _sid(thread_id)})
    except gateway.AuthRequired:
        return {"ok": False, "error": "auth_required"}
    except gateway.GatewayError as e:
        return {"ok": False, "error": str(e)}
    indexed = res.get("indexed", len(body.documents))
    chars = sum(len(d) for d in body.documents)
    aid = db.run("INSERT INTO attachments(thread_id,name,chars,chunks,created_at) VALUES(?,?,?,?,?)",
                 (thread_id, body.name, chars, indexed, db.now()))
    return {"ok": True, "id": aid, "indexed": indexed, "name": body.name}


@router.get("/threads/{thread_id}/files")
def files(thread_id: int) -> list[dict]:
    """Проиндексированные вложения треда — чтобы видеть, что уже в базе знаний."""
    return db.q("SELECT id,name,chars,chunks,created_at FROM attachments WHERE thread_id=? ORDER BY id DESC",
                (thread_id,))


@router.delete("/threads/{thread_id}/files/{att_id}")
def del_file(thread_id: int, att_id: int) -> dict:
    # из локального списка убираем; из Qdrant чанки живут по TTL сессии (чистится шлюзом)
    db.run("DELETE FROM attachments WHERE id=? AND thread_id=?", (att_id, thread_id))
    return {"ok": True}


@router.get("/agent-roles")
def agent_roles() -> list[dict]:
    return [{"id": k, "name": v[0], "brief": v[1], "default": k in DEFAULT_ROLES}
            for k, v in ROLE_PRESETS.items()]


# ── экспорт треда в файл в «Загрузки» (md/docx/xlsx; pdf делает Electron) ──
@router.post("/threads/{thread_id}/export")
def export_thread(thread_id: int, body: ExportIn) -> dict:
    th = db.q("SELECT title FROM threads WHERE id=?", (thread_id,))
    if not th:
        return {"ok": False, "error": "no_thread"}
    title = th[0]["title"] or "chat"
    msgs = db.q("SELECT role,content,meta FROM messages WHERE thread_id=? ORDER BY id", (thread_id,))
    base, out, fmt = _safe(title), _downloads(), body.format.lower()
    try:
        if fmt == "md":
            p = out / (base + ".md")
            lines = [f"# {title}\n"]
            for m in msgs:
                who = "🧑 Вы" if m["role"] == "user" else "🤖 Ассистент"
                lines.append(f"\n## {who}\n\n{m['content']}\n")
            p.write_text("\n".join(lines), encoding="utf-8")
        elif fmt == "docx":
            from docx import Document
            doc = Document()
            doc.add_heading(title, 0)
            for m in msgs:
                doc.add_heading("Вы" if m["role"] == "user" else "Ассистент", level=2)
                doc.add_paragraph(m["content"])
            p = out / (base + ".docx")
            doc.save(str(p))
        elif fmt == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "chat"
            ws.append(["Роль", "Сообщение", "Модель", "Стоимость ₽"])
            for m in msgs:
                meta = json.loads(m["meta"] or "{}")
                ws.append([m["role"], m["content"], meta.get("model", ""), meta.get("cost_rub", "")])
            p = out / (base + ".xlsx")
            wb.save(str(p))
        else:
            return {"ok": False, "error": "bad_format"}
    except Exception as e:  # noqa: BLE001
        return {"ok": False, "error": str(e)}
    return {"ok": True, "path": str(p)}


# ── отправка сообщения ──
@router.post("/threads/{thread_id}/send")
def send(thread_id: int, body: SendIn) -> dict:
    th = db.q("SELECT profile,skills FROM threads WHERE id=?", (thread_id,))
    if not th:
        return {"ok": False, "error": "no_thread"}
    profile = th[0]["profile"] or "standard"
    skills = [s for s in (th[0]["skills"] or "").split(",") if s]
    db.run("INSERT INTO messages(thread_id,role,content,meta,created_at) VALUES(?,?,?,?,?)",
           (thread_id, "user", body.prompt, "{}", db.now()))

    # контекст из вложений (best-effort)
    ctx = ""
    try:
        found = gateway.call("rag_search", {"query": body.prompt, "session_id": _sid(thread_id), "top_k": 3})
        chunks = [r.get("text", "") for r in (found.get("results") or [])]
        if chunks:
            ctx = "Контекст из приложенных документов:\n- " + "\n- ".join(chunks) + "\n\n"
    except (gateway.GatewayError, Exception):  # noqa: BLE001 — RAG опционален
        pass

    hist = _history(thread_id)
    prompt = ctx + (f"История:\n{hist}\n\n" if hist else "") + f"Ты: {body.prompt}"
    try:
        res = gateway.call("chat", {"prompt": prompt, "session_id": _sid(thread_id),
                                    "profile": profile, "system": _system_for(skills), "max_tokens": 1500})
    except gateway.AuthRequired:
        return {"ok": False, "error": "auth_required"}
    except gateway.GatewayError as e:
        return {"ok": False, "error": str(e)}

    text = res.get("text", "")
    meta = {"model": res.get("model"), "cost_rub": res.get("cost_rub"),
            "input_tokens": res.get("input_tokens"), "output_tokens": res.get("output_tokens")}
    mid = db.run("INSERT INTO messages(thread_id,role,content,meta,created_at) VALUES(?,?,?,?,?)",
                 (thread_id, "assistant", text, json.dumps(meta, ensure_ascii=False), db.now()))
    db.run("UPDATE threads SET updated_at=? WHERE id=?", (db.now(), thread_id))
    return {"ok": True, "id": mid, "content": text, "meta": meta}


def _sse(d: dict) -> str:
    return "data: " + json.dumps(d, ensure_ascii=False) + "\n\n"


def _build_prompt(thread_id: int, user_prompt: str) -> str:
    """Контекст из вложений (если есть) + история + текущий вопрос."""
    ctx = ""
    if db.q("SELECT 1 FROM attachments WHERE thread_id=? LIMIT 1", (thread_id,)):
        try:
            found = gateway.call("rag_search", {"query": user_prompt, "session_id": _sid(thread_id), "top_k": 3})
            chunks = [r.get("text", "") for r in (found.get("results") or [])]
            if chunks:
                ctx = "Контекст из приложенных документов:\n- " + "\n- ".join(chunks) + "\n\n"
        except Exception:  # noqa: BLE001
            pass
    hist = _history(thread_id)
    return ctx + (f"История:\n{hist}\n\n" if hist else "") + f"Ты: {user_prompt}"


# ── стриминг: ответ появляется постепенно (SSE от portal_api → UI) ──
@router.post("/threads/{thread_id}/send-stream")
def send_stream(thread_id: int, body: SendIn) -> StreamingResponse:
    th = db.q("SELECT profile,skills FROM threads WHERE id=?", (thread_id,))

    def gen():
        if not th:
            yield _sse({"error": "no_thread"}); return
        tok = auth.token()
        if not tok:
            yield _sse({"error": "auth_required"}); return
        profile = th[0]["profile"] or "standard"
        skills = [s for s in (th[0]["skills"] or "").split(",") if s]
        db.run("INSERT INTO messages(thread_id,role,content,meta,created_at) VALUES(?,?,?,?,?)",
               (thread_id, "user", body.prompt, "{}", db.now()))
        payload = {"prompt": _build_prompt(thread_id, body.prompt), "session_id": _sid(thread_id),
                   "profile": profile, "system": _system_for(skills), "max_tokens": 1500}
        req = urllib.request.Request(
            config.PORTAL + "/api/chat/stream", data=json.dumps(payload).encode(),
            headers={"Authorization": "Bearer " + tok, "Content-Type": "application/json"})
        full, meta = "", {}
        try:
            r = urllib.request.urlopen(req, timeout=300)
            for raw in r:
                line = raw.decode("utf-8", "replace").strip()
                if not line.startswith("data:"):
                    continue
                try:
                    d = json.loads(line[5:].strip())
                except Exception:  # noqa: BLE001
                    continue
                if d.get("delta"):
                    full += d["delta"]
                    yield _sse({"delta": d["delta"]})
                elif d.get("done"):
                    meta = {"model": d.get("model"), "cost_rub": d.get("cost_rub"),
                            "input_tokens": d.get("input_tokens"), "output_tokens": d.get("output_tokens")}
                elif d.get("error"):
                    yield _sse({"error": d["error"]})
        except urllib.error.HTTPError as e:
            yield _sse({"error": f"HTTP {e.code}"})
        except Exception as e:  # noqa: BLE001
            yield _sse({"error": str(e)})
        if full:
            db.run("INSERT INTO messages(thread_id,role,content,meta,created_at) VALUES(?,?,?,?,?)",
                   (thread_id, "assistant", full, json.dumps(meta, ensure_ascii=False), db.now()))
            db.run("UPDATE threads SET updated_at=? WHERE id=?", (db.now(), thread_id))
        yield _sse({"done": True, "meta": meta})

    return StreamingResponse(gen(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache"})


def _catalog_system(a: dict) -> str:
    """system-prompt агента из каталога (методика в полях)."""
    parts = [f"Ты — {a['name']}."]
    if a.get("description"):
        parts.append(a["description"])
    if a.get("steps"):
        parts.append("Методика (шаги):\n" + a["steps"])
    if a.get("dod"):
        parts.append("Definition of Done:\n" + a["dod"])
    if a.get("antipatterns"):
        parts.append("Избегай (анти-паттерны):\n" + a["antipatterns"])
    hints = [f"[{s}] {SKILLS[s]}" for s in (a.get("skills") or "").split(",") if s in SKILLS]
    if hints:
        parts.append("Методики-скиллы:\n" + "\n".join(hints))
    return "\n\n".join(parts)


# ── мультиагенты: каталог (agent_ids) ИЛИ быстрые роли (roles), цепочкой в текущий тред ──
@router.post("/threads/{thread_id}/agents")
def agents(thread_id: int, body: AgentsIn) -> dict:
    # specs: список (имя, system) — из каталога приоритетно, иначе из пресетов
    specs = []
    if body.agent_ids:
        for aid in body.agent_ids:
            rows = db.q("SELECT name,description,skills,steps,dod,antipatterns FROM agents WHERE id=?", (aid,))
            if rows:
                specs.append((rows[0]["name"], _catalog_system(rows[0])))
    if not specs:
        roles = [r for r in (body.roles or DEFAULT_ROLES) if r in ROLE_PRESETS] or DEFAULT_ROLES
        specs = [(ROLE_PRESETS[r][0], f"Ты — {ROLE_PRESETS[r][0]}. {ROLE_PRESETS[r][1]}") for r in roles]
    names = ", ".join(n for n, _ in specs)
    db.run("INSERT INTO messages(thread_id,role,content,meta,created_at) VALUES(?,?,?,?,?)",
           (thread_id, "user", f"[агенты: {names}] {body.task}", "{}", db.now()))
    outputs, prior = [], ""
    try:
        for name, system in specs:
            prefix = ("Наработки предыдущих ролей:\n" + prior) if prior else ""
            r = gateway.call("chat", {"prompt": f"Задача: {body.task}\n\n{prefix}",
                                      "session_id": _sid(thread_id) + "-agents", "profile": "standard",
                                      "system": system, "max_tokens": 1200})
            t = r.get("text", "")
            outputs.append(f"### {name}\n{t}")
            prior += f"\n[{name}]: {t}\n"
    except gateway.AuthRequired:
        return {"ok": False, "error": "auth_required"}
    except gateway.GatewayError as e:
        return {"ok": False, "error": str(e)}
    combined = "\n\n".join(outputs)
    mid = db.run("INSERT INTO messages(thread_id,role,content,meta,created_at) VALUES(?,?,?,?,?)",
                 (thread_id, "assistant", combined, json.dumps({"agents": names}), db.now()))
    db.run("UPDATE threads SET updated_at=? WHERE id=?", (db.now(), thread_id))
    return {"ok": True, "id": mid, "content": combined}
